from django.shortcuts import render
from django.contrib.auth.decorators import login_required, user_passes_test
from django.http import HttpResponse
from api_client import EmulatorAPI
from accounts.models import Profile
from collections import OrderedDict
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from io import BytesIO

def is_dean(user):
    try:
        return user.profile.role == 'dean'
    except Profile.DoesNotExist:
        return False

def get_all_data(api):
    all_debts = api.get_debts()
    disciplines = api.get_disciplines()
    teachers = api.get_teachers()
    groups = api.get_groups()
    accounts = api.get_accounts()
    return all_debts, disciplines, teachers, groups, accounts

def enrich_debts(all_debts, disciplines, teachers, groups, accounts):
    disc_map = {d['id']: d for d in disciplines}
    
    teacher_map = {}
    for t in teachers:
        parts = [t.get('last_name', ''), t.get('first_name', ''), t.get('middle_name', '')]
        teacher_map[t['id']] = {
            'full_name': ' '.join(parts).strip(),
            'department': t.get('department', '')
        }
    
    group_map = {g['id']: g.get('name', '—') for g in groups}

    student_map = {}
    for acc in accounts:
        if acc.get('linked_entity_type') == 'student' and acc.get('linked_entity_id'):
            sid = acc['linked_entity_id']
            student_map[sid] = {
                'email': acc.get('email', '—'),
                'full_name': f"{acc.get('last_name', '')} {acc.get('first_name', '')} {acc.get('middle_name', '')}".strip(),
                'group_id': acc.get('group_id', None),
            }

    CONTROL_TYPES = {
        'exam': 'Экзамен',
        'credit': 'Зачёт',
        'differentiated_credit': 'Диф. зачёт',
        'coursework': 'Курсовая',
        'practice': 'Практика',
        'lab_work': 'Лаб. работа',
        'other': 'Другое',
    }

    DEBT_TYPES = {
        'failed_exam': 'Не сдан экзамен',
        'failed_credit': 'Не сдан зачёт',
        'not_attended': 'Неявка',
        'not_allowed': 'Не допущен',
        'coursework_not_submitted': 'Курсовая не сдана',
        'practice_not_completed': 'Практика не пройдена',
        'other': 'Другая причина',
    }

    STATUS_MAP = {
        'open': 'Открыт',
        'in_progress': 'В процессе',
        'passed': 'Сдано',
        'failed': 'Задолженность',
        'closed': 'Сдано',
    }

    for debt in all_debts:
        disc_id = debt.get('discipline_id')
        disc = disc_map.get(disc_id, {})
        
        student_id = debt.get('student_id')
        student = student_map.get(student_id, {})
        
        teacher = teacher_map.get(debt.get('teacher_id'), {})
        
        debt['discipline_name'] = disc.get('name', '—')
        debt['discipline_code'] = disc.get('code', '')
        debt['department'] = disc.get('department', '')
        debt['student_name'] = student.get('full_name', '—')
        debt['student_email'] = student.get('email', '—')
        debt['group_name'] = group_map.get(student.get('group_id'), '—')
        debt['teacher_name'] = teacher.get('full_name', '—')
        debt['teacher_department'] = teacher.get('department', '')
        debt['control_type_display'] = CONTROL_TYPES.get(debt.get('control_type', ''), debt.get('control_type', '—'))
        debt['debt_type_display'] = DEBT_TYPES.get(debt.get('debt_type', ''), debt.get('debt_type', '—'))
        debt['status_display'] = STATUS_MAP.get(debt.get('status', ''), debt.get('status', '—'))
        debt['semester_display'] = f"{debt.get('semester', '?')} семестр"

    return all_debts

def apply_filters(debts, search='', group='', status='', control=''):
    filtered = []
    for d in debts:
        if search:
            search_lower = search.lower()
            if (search_lower not in d.get('student_name', '').lower() and
                search_lower not in d.get('student_email', '').lower() and
                search_lower not in d.get('group_name', '').lower() and
                search_lower not in d.get('discipline_name', '').lower()):
                continue
        
        if group and d.get('group_name', '') != group:
            continue
        
        if status:
            if status == 'passed':
                if d.get('status') not in ('passed', 'closed'):
                    continue
            else:
                if d.get('status') != status:
                    continue
        
        if control and d.get('control_type', '') != control:
            continue
        
        filtered.append(d)
    
    return filtered

def group_by_discipline(debts):
    grouped = OrderedDict()
    for d in debts:
        disc_id = d.get('discipline_id')
        if disc_id not in grouped:
            grouped[disc_id] = {
                'discipline_name': d.get('discipline_name', '—'),
                'discipline_code': d.get('discipline_code', ''),
                'department': d.get('department', ''),
                'teacher_name': d.get('teacher_name', '—'),
                'debts': [],
                'total_count': 0,
                'open_count': 0,
                'passed_count': 0,
                'failed_count': 0,
                'other_count': 0,
            }
        
        grouped[disc_id]['debts'].append(d)
        grouped[disc_id]['total_count'] += 1
        
        st = d.get('status', '')
        if st in ['open', 'in_progress']:
            grouped[disc_id]['open_count'] += 1
        elif st in ('passed', 'closed'):
            grouped[disc_id]['passed_count'] += 1
        elif st == 'failed':
            grouped[disc_id]['failed_count'] += 1
        else:
            grouped[disc_id]['other_count'] += 1
    
    return grouped

@login_required
@user_passes_test(is_dean)
def debts_summary(request):
    api = EmulatorAPI()
    
    search = request.GET.get('search', '').strip()
    group = request.GET.get('group', '').strip()
    status = request.GET.get('status', '').strip()
    control = request.GET.get('control', '').strip()
    
    try:
        all_debts, disciplines, teachers, groups, accounts = get_all_data(api)
        enriched = enrich_debts(all_debts, disciplines, teachers, groups, accounts)
        filtered = apply_filters(enriched, search, group, status, control)
        grouped = group_by_discipline(filtered)
    except Exception as e:
        print(f"Error: {e}")
        return render(request, 'reports/debts_summary.html', {
            'grouped_summary': {},
            'total_debts': 0,
            'groups': [],
            'control_types': [],
            'error': str(e)
        })

    all_enriched = enrich_debts(all_debts, disciplines, teachers, groups, accounts)
    all_groups = sorted({d.get('group_name') for d in all_enriched if d.get('group_name') != '—'})
    
    CONTROL_TYPES_FILTER = {
        'exam': 'Экзамен',
        'credit': 'Зачёт',
        'differentiated_credit': 'Диф. зачёт',
        'coursework': 'Курсовая',
        'practice': 'Практика',
        'lab_work': 'Лаб. работа',
        'other': 'Другое',
    }
    raw_controls = sorted({d.get('control_type') for d in all_enriched if d.get('control_type')})
    all_controls = [(c, CONTROL_TYPES_FILTER.get(c, c)) for c in raw_controls]

    request.session['report_filters'] = {
        'search': search,
        'group': group,
        'status': status,
        'control': control,
    }

    return render(request, 'reports/debts_summary.html', {
        'grouped_summary': grouped,
        'total_debts': len(filtered),
        'groups': all_groups,
        'control_types': all_controls,
        'current_filters': {
            'search': search,
            'group': group,
            'status': status,
            'control': control,
        }
    })

@login_required
@user_passes_test(is_dean)
def export_debts_excel(request):
    api = EmulatorAPI()
    
    filters = {
        'search': request.GET.get('search', '').strip(),
        'group': request.GET.get('group', '').strip(),
        'status': request.GET.get('status', '').strip(),
        'control': request.GET.get('control', '').strip(),
    }
    
    all_debts, disciplines, teachers, groups, accounts = get_all_data(api)
    enriched = enrich_debts(all_debts, disciplines, teachers, groups, accounts)
    filtered = apply_filters(
        enriched,
        filters['search'],
        filters['group'],
        filters['status'],
        filters['control']
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Задолженности"

    headers = ['Дисциплина', 'Студент', 'Email', 'Группа', 'Преподаватель', 'Тип контроля', 'Причина', 'Семестр', 'Статус', 'Оценка']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for row_idx, d in enumerate(filtered, 2):
        ws.cell(row=row_idx, column=1, value=d.get('discipline_name', '—'))
        ws.cell(row=row_idx, column=2, value=d.get('student_name', '—'))
        ws.cell(row=row_idx, column=3, value=d.get('student_email', '—'))
        ws.cell(row=row_idx, column=4, value=d.get('group_name', '—'))
        ws.cell(row=row_idx, column=5, value=d.get('teacher_name', '—'))
        ws.cell(row=row_idx, column=6, value=d.get('control_type_display', '—'))
        ws.cell(row=row_idx, column=7, value=d.get('debt_type_display', '—'))
        ws.cell(row=row_idx, column=8, value=d.get('semester_display', '—'))
        ws.cell(row=row_idx, column=9, value=d.get('status_display', '—'))
        
        grade = d.get('grade')
        if isinstance(grade, dict):
            if grade.get('type') == 'numeric':
                grade_str = f"{grade.get('value', '')} ({grade.get('label', '')})"
            elif grade.get('type') == 'pass_fail':
                grade_str = 'Зачёт' if grade.get('value') == 'passed' else 'Незачёт' if grade.get('value') == 'failed' else str(grade.get('value', ''))
            else:
                grade_str = str(grade)
        else:
            grade_str = str(grade) if grade else '—'
        ws.cell(row=row_idx, column=10, value=grade_str)

    for col in range(1, 11):
        ws.column_dimensions[get_column_letter(col)].width = 20

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="debts_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
    return response

@login_required
@user_passes_test(is_dean)
def export_debts_word(request):
    api = EmulatorAPI()
    
    filters = {
        'search': request.GET.get('search', '').strip(),
        'group': request.GET.get('group', '').strip(),
        'status': request.GET.get('status', '').strip(),
        'control': request.GET.get('control', '').strip(),
    }
    
    all_debts, disciplines, teachers, groups, accounts = get_all_data(api)
    enriched = enrich_debts(all_debts, disciplines, teachers, groups, accounts)
    filtered = apply_filters(
        enriched,
        filters['search'],
        filters['group'],
        filters['status'],
        filters['control']
    )

    total = len(filtered)
    open_count = sum(1 for d in filtered if d.get('status') in ['open', 'in_progress'])
    passed_count = sum(1 for d in filtered if d.get('status') in ['passed', 'closed'])
    failed_count = sum(1 for d in filtered if d.get('status') == 'failed')
    other_count = total - open_count - passed_count - failed_count

    grouped = group_by_discipline(filtered)

    # Улучшенные стили для Word – таблица вписывается в страницу
    html = f'''
    <html xmlns:o="urn:schemas-microsoft-com:office:office"
          xmlns:w="urn:schemas-microsoft-com:office:word"
          xmlns="http://www.w3.org/TR/REC-html40">
    <head>
        <meta charset="utf-8">
        <title>Краткий отчёт по задолженностям</title>
        <style>
            body {{
                font-family: 'Times New Roman', serif;
                font-size: 12pt;
                line-height: 1.3;
                margin: 2cm auto;
                width: 90%;
                max-width: 1000px;
            }}
            h1 {{ text-align: center; font-size: 18pt; margin-bottom: 20px; }}
            h2 {{ font-size: 16pt; margin-top: 20px; margin-bottom: 10px; }}
            p {{ margin: 10px 0; }}
            .bad {{ color: #c0392b; font-weight: bold; }}
            .good {{ color: #27ae60; font-weight: bold; }}
            .warn {{ color: #f39c12; font-weight: bold; }}
            .date {{ text-align: right; color: #7f8c8d; margin-bottom: 30px; }}
            table {{
                width: 100%;
                border-collapse: collapse;
                table-layout: fixed;
                font-size: 10pt;
                margin: 15px 0;
            }}
            th, td {{
                border: 1px solid #aaa;
                padding: 5px 4px;
                vertical-align: top;
                word-wrap: break-word;
            }}
            th {{
                background-color: #667eea;
                color: white;
                font-weight: bold;
            }}
            /* Ширина колонок в процентах */
            .summary-table th:nth-child(1) {{ width: 30%; }}
            .summary-table th:nth-child(2) {{ width: 20%; }}
            .summary-table th:nth-child(3) {{ width: 10%; }}
            .summary-table th:nth-child(4) {{ width: 10%; }}
            .summary-table th:nth-child(5) {{ width: 10%; }}
            .summary-table th:nth-child(6) {{ width: 10%; }}
            .summary-table th:nth-child(7) {{ width: 10%; }}
            
            .debt-table th:nth-child(1) {{ width: 25%; }}
            .debt-table th:nth-child(2) {{ width: 25%; }}
            .debt-table th:nth-child(3) {{ width: 15%; }}
            .debt-table th:nth-child(4) {{ width: 15%; }}
            .debt-table th:nth-child(5) {{ width: 10%; }}
            .debt-table th:nth-child(6) {{ width: 10%; }}
            
            @media print {{
                body {{ margin: 0; }}
                table {{ page-break-inside: avoid; }}
            }}
        </style>
    </head>
    <body>
        <h1>Краткий отчёт по академическим задолженностям</h1>
        <p class="date">Дата: {datetime.now().strftime("%d.%m.%Y %H:%M")}</p>
        
        <h2>1. Общая сводка</h2>
        <p>Всего задолженностей: <span class="bad">{total}</span>.</p>
        <p>В процессе: <span class="warn">{open_count}</span> | Сдано: <span class="good">{passed_count}</span> | Просрочено: <span class="bad">{failed_count}</span> | Прочее: {other_count}</p>
        
        <h2>2. Сводка по дисциплинам</h2>
        <table class="summary-table">
            <thead>
                <tr><th>Дисциплина</th><th>Преподаватель</th><th>Всего</th><th>Открыто</th><th>Сдано</th><th>Просрочено</th><th>Прочее</th></tr>
            </thead>
            <tbody>
    '''
    for disc_id, data in grouped.items():
        html += f'''
        <tr>
            <td>{data['discipline_name']}</td>
            <td>{data['teacher_name']}</td>
            <td style="text-align:center">{data['total_count']}</td>
            <td style="text-align:center">{data['open_count']}</td>
            <td style="text-align:center">{data['passed_count']}</td>
            <td style="text-align:center">{data['failed_count']}</td>
            <td style="text-align:center">{data['other_count']}</td>
        </tr>
        '''
    html += '''
            </tbody>
        </table>
    '''
    
    if grouped:
        html += '<h2>3. Детализация задолженностей (первые 50 записей)</h2>'
        html += '''
        <table class="debt-table">
            <thead>
                <tr><th>Студент</th><th>Группа</th><th>Дисциплина</th><th>Тип контроля</th><th>Статус</th><th>Оценка</th></tr>
            </thead>
            <tbody>
        '''
        shown = 0
        for disc_id, data in grouped.items():
            for d in data['debts']:
                if shown >= 50:
                    break
                grade_val = ''
                if d.get('grade'):
                    if isinstance(d['grade'], dict):
                        if d['grade'].get('type') == 'numeric':
                            grade_val = f"{d['grade'].get('value', '')}"
                        elif d['grade'].get('type') == 'pass_fail':
                            grade_val = 'Зачёт' if d['grade'].get('value') == 'passed' else 'Незачёт'
                    else:
                        grade_val = str(d['grade'])
                else:
                    grade_val = '—'
                status_val = d.get('status_display', '—')
                html += f'''
                <tr>
                    <td>{d.get('student_name', '—')}</td>
                    <td>{d.get('group_name', '—')}</td>
                    <td>{d.get('discipline_name', '—')}</td>
                    <td>{d.get('control_type_display', '—')}</td>
                    <td>{status_val}</td>
                    <td>{grade_val}</td>
                </tr>
                '''
                shown += 1
            if shown >= 50:
                break
        html += '''
            </tbody>
        </table>
        <p><em>Примечание: показаны первые 50 записей. Полный список доступен в Excel-отчёте.</em></p>
        '''
    
    html += '''
        <br>
        <p>Отчёт подготовлен автоматически.</p>
    </body>
    </html>
    '''
    response = HttpResponse(html, content_type='application/msword; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.doc"'
    return response